# -*- coding: utf-8 -*-
"""PPT 几何质量检查：越界 / 图文重叠 / 文本溢出 估计。输出问题清单。"""
import math
import sys
from pptx import Presentation

EMU = 914400.0
SW, SH = 13.333, 7.5

path = sys.argv[1] if len(sys.argv) > 1 else "docs/HRMS答辩汇报.pptx"
prs = Presentation(path)
issues = []


def I(v):
    return v / EMU


def est_text_height(sh, w_in):
    """估算文本渲染高度(pt)：CJK≈字号宽，ASCII≈0.55字号；行高1.22。"""
    total = 0.0
    tf = sh.text_frame
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        if not text:
            continue
        size = None
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        size = size or 14.0
        cjk = sum(1 for c in text if ord(c) > 0x2E80)
        asc = len(text) - cjk
        text_pt = cjk * size + asc * size * 0.55
        line_w = w_in * 72.0
        lines = max(1, math.ceil(text_pt / max(line_w, 1)))
        sa = p.space_after.pt if p.space_after else 0
        total += lines * size * 1.22 + sa
    return total


for si, slide in enumerate(prs.slides, 1):
    pics, texts = [], []
    for sh in slide.shapes:
        if sh.left is None:
            continue
        x, y, w, h = I(sh.left), I(sh.top), I(sh.width), I(sh.height)
        if x < -0.02 or y < -0.02 or x + w > SW + 0.02 or y + h > SH + 0.02:
            issues.append(
                f"S{si:02d} 越界: {sh.shape_type} ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f} "
                f"txt={sh.text_frame.text[:12] if sh.has_text_frame else ''}")
        if str(sh.shape_type) == "PICTURE (13)":
            pics.append((x, y, w, h))
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append((sh, x, y, w, h))
    # 图片与文本框重叠
    for px, py, pw, ph in pics:
        for sh, tx, ty, tw, th in texts:
            ox = max(0, min(px + pw, tx + tw) - max(px, tx))
            oy = max(0, min(py + ph, ty + th) - max(py, ty))
            if ox > 0.05 and oy > 0.05:
                issues.append(
                    f"S{si:02d} 图文重叠: 图片({px:.2f},{py:.2f},{pw:.2f}x{ph:.2f}) ×"
                    f"文本({tx:.2f},{ty:.2f}) '{sh.text_frame.text[:14]}'")
    # 文本溢出
    for sh, tx, ty, tw, th in texts:
        est = est_text_height(sh, tw)
        if est > th * 72 * 1.15 + 3:
            issues.append(
                f"S{si:02d} 文本溢出: '{sh.text_frame.text[:16]}…' 框高{th*72:.0f}pt "
                f"估算{est:.0f}pt ({tx:.2f},{ty:.2f},{tw:.2f}x{th:.2f})")

print(f"检查完成：{len(prs.slides._sldIdLst)} 页，发现 {len(issues)} 个问题")
for it in issues:
    print(" -", it)
